"""Turn any report table into a CSV, Excel or PDF download.

A report describes itself once as a :class:`Sheet` — title, the period and
filters it was run for, columns, rows and totals — and :func:`file_response`
writes that same sheet in whichever format was asked for. The figures in the
file are therefore exactly the figures on screen, in all three formats.

Rows are dicts keyed by column. A row may carry ``"_style"`` to set it apart:
``heading`` (an account or section title), ``opening`` / ``closing`` (a
ledger's carried balances), ``subtotal`` and ``total``.
"""
from __future__ import annotations

import csv
import datetime as _dt
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone

FORMATS = {
    "csv": ("text/csv; charset=utf-8", "csv"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "pdf": ("application/pdf", "pdf"),
}

COMPANY = "Wilkem Ventures"


@dataclass
class Column:
    key: str
    label: str
    #: text | money | date | number
    kind: str = "text"
    #: Relative width hint for PDF and Excel.
    width: int = 12


@dataclass
class Sheet:
    title: str
    columns: list[Column]
    rows: list[dict]
    #: Lines printed under the title: period, building, filters.
    subtitle: list[str] = field(default_factory=list)
    #: Short lines printed after the table (a basis note, a balance check).
    notes: list[str] = field(default_factory=list)


def _plain(value, kind: str):
    """A cell's value as a CSV / Excel cell wants it."""
    if value is None or value == "":
        return ""
    if kind == "money":
        return Decimal(str(value)).quantize(Decimal("0.01"))
    if kind == "date" and isinstance(value, str):
        try:
            return _dt.date.fromisoformat(value)
        except ValueError:
            return value
    return value


def _printed(value, kind: str) -> str:
    """A cell as it reads on paper."""
    if value is None or value == "":
        return ""
    if kind == "money":
        amount = Decimal(str(value)).quantize(Decimal("0.01"))
        text = f"{abs(amount):,.2f}"
        return f"({text})" if amount < 0 else text
    if kind == "date":
        day = value if isinstance(value, _dt.date) else _dt.date.fromisoformat(str(value))
        return day.strftime("%d %b %Y")
    return str(value)


def _stamp() -> str:
    return timezone.localtime().strftime("%d %b %Y %H:%M")


# ── CSV ──────────────────────────────────────────────────────────────────────

def to_csv(sheet: Sheet) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow([f"{COMPANY} — {sheet.title}"])
    for line in sheet.subtitle:
        writer.writerow([line])
    writer.writerow([f"Printed {_stamp()}"])
    writer.writerow([])
    writer.writerow([c.label for c in sheet.columns])
    for row in sheet.rows:
        cells = []
        for c in sheet.columns:
            value = _plain(row.get(c.key), c.kind)
            if isinstance(value, _dt.date):
                value = value.isoformat()
            cells.append(value)
        writer.writerow(cells)
    if sheet.notes:
        writer.writerow([])
        for note in sheet.notes:
            writer.writerow([note])
    # BOM so Excel opens the UTF-8 (KES, em dashes, names) correctly.
    return ("﻿" + out.getvalue()).encode("utf-8")


# ── Excel ────────────────────────────────────────────────────────────────────

def to_xlsx(sheet: Sheet) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "", sheet.title)[:31] or "Report"

    ncols = len(sheet.columns)
    ws.cell(row=1, column=1, value=f"{COMPANY} — {sheet.title}").font = Font(bold=True, size=14)
    r = 2
    for line in [*sheet.subtitle, f"Printed {_stamp()}"]:
        ws.cell(row=r, column=1, value=line).font = Font(color="636776", size=10)
        r += 1
    r += 1

    header_row = r
    header_fill = PatternFill("solid", fgColor="F0EDE5")
    for i, col in enumerate(sheet.columns, start=1):
        cell = ws.cell(row=r, column=i, value=col.label)
        cell.font = Font(bold=True, color="404554")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="right" if col.kind in ("money", "number") else "left")
    r += 1

    thin = Side(style="thin", color="999999")
    styles = {
        "heading": (Font(bold=True, size=11), PatternFill("solid", fgColor="F7F5F0"), None),
        "opening": (Font(italic=True, color="555555"), None, None),
        "closing": (Font(bold=True), None, Border(top=thin)),
        "subtotal": (Font(bold=True), None, Border(top=thin)),
        "total": (Font(bold=True), PatternFill("solid", fgColor="F0EDE5"), Border(top=thin, bottom=thin)),
    }
    for row in sheet.rows:
        font, fill, border = styles.get(row.get("_style"), (None, None, None))
        for i, col in enumerate(sheet.columns, start=1):
            value = _plain(row.get(col.key), col.kind)
            if isinstance(value, Decimal):
                value = float(value)
            cell = ws.cell(row=r, column=i, value=value if value != "" else None)
            if col.kind == "money":
                cell.number_format = '#,##0.00;(#,##0.00);"-"'
            elif col.kind == "date" and isinstance(value, _dt.date):
                cell.number_format = "dd mmm yyyy"
            if font:
                cell.font = font
            if fill:
                cell.fill = fill
            if border:
                cell.border = border
        r += 1

    if sheet.notes:
        r += 1
        for note in sheet.notes:
            ws.cell(row=r, column=1, value=note).font = Font(italic=True, color="636776", size=9)
            r += 1

    for i, col in enumerate(sheet.columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(col.width, len(col.label) + 2)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{max(header_row, r - 1)}"
    ws.page_setup.orientation = "landscape" if ncols > 5 else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = f"{header_row}:{header_row}"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────────────

def to_pdf(sheet: Sheet) -> bytes:
    """A4 PDF drawn with ReportLab's table layout.

    Not the HTML route the statements use: xhtml2pdf took over half a minute
    on a nine-month general ledger, past the server's request timeout, where
    ReportLab draws the same thousand-line table in a second or two.
    """
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    wide = len(sheet.columns) > 5
    pagesize = landscape(A4) if wide else A4
    margin = 1.2 * cm
    usable = pagesize[0] - 2 * margin

    base = ParagraphStyle("cell", fontName="Helvetica", fontSize=7.5, leading=9)
    bold = ParagraphStyle("cellb", parent=base, fontName="Helvetica-Bold")
    head = ParagraphStyle("head", parent=base, fontName="Helvetica-Bold", fontSize=7,
                          textColor=colors.HexColor("#404554"))
    title_style = ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=15, leading=18)
    brand_style = ParagraphStyle("brand", fontName="Helvetica", fontSize=8,
                                 textColor=colors.HexColor("#636776"), leading=10)
    sub_style = ParagraphStyle("sub", fontName="Helvetica", fontSize=8.5, leading=11,
                               textColor=colors.HexColor("#444444"))
    note_style = ParagraphStyle("note", parent=sub_style, fontSize=7.5, leading=10)

    total_width = sum(c.width for c in sheet.columns)
    widths = [usable * c.width / total_width for c in sheet.columns]
    numeric = [c.kind in ("money", "number") for c in sheet.columns]
    # Long text wraps; short columns stay plain strings, which draw much faster.
    wraps = [not num and w > 3.2 * cm for num, w in zip(numeric, widths, strict=True)]

    head_right = ParagraphStyle("headr", parent=head, alignment=2)
    data = [[Paragraph(escape(c.label.upper()), head_right if num else head)
             for c, num in zip(sheet.columns, numeric, strict=True)]]
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F0EDE5")),
        ("LINEBELOW", (0, 0), (-1, 0), 0.75, colors.HexColor("#999999")),
        ("FONT", (0, 1), (-1, -1), "Helvetica", 7.5),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#E1E1E6")),
    ]
    for i, num in enumerate(numeric):
        if num:
            commands.append(("ALIGN", (i, 0), (i, -1), "RIGHT"))

    for r, row in enumerate(sheet.rows, start=1):
        style = row.get("_style", "")
        strong = style in ("heading", "closing", "subtotal", "total")
        cells = []
        for i, col in enumerate(sheet.columns):
            text = _printed(row.get(col.key), col.kind)
            cells.append(Paragraph(escape(text), bold if strong else base) if wraps[i] and text else text)
        data.append(cells)
        if strong:
            commands.append(("FONT", (0, r), (-1, r), "Helvetica-Bold", 7.5))
        if style == "heading":
            commands.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#F7F5F0")))
        elif style == "opening":
            commands.append(("FONT", (0, r), (-1, r), "Helvetica-Oblique", 7.5))
            commands.append(("TEXTCOLOR", (0, r), (-1, r), colors.HexColor("#555555")))
        elif style in ("closing", "subtotal"):
            commands.append(("LINEABOVE", (0, r), (-1, r), 0.6, colors.HexColor("#999999")))
        elif style == "total":
            commands.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#F0EDE5")))
            commands.append(("LINEABOVE", (0, r), (-1, r), 0.9, colors.HexColor("#555555")))
            commands.append(("LINEBELOW", (0, r), (-1, r), 0.9, colors.HexColor("#555555")))
    if not sheet.rows:
        data.append(["Nothing to report for this period."] + [""] * (len(sheet.columns) - 1))

    table = LongTable(data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle(commands))

    story = [Paragraph(escape(COMPANY.upper()), brand_style), Paragraph(escape(sheet.title), title_style)]
    story += [Paragraph(escape(line), sub_style) for line in sheet.subtitle]
    story += [Spacer(1, 8), table]
    if sheet.notes:
        story.append(Spacer(1, 8))
        story += [Paragraph(escape(n), note_style) for n in sheet.notes]

    printed = _stamp()

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawString(margin, 0.7 * cm, f"{COMPANY} · {sheet.title} · Printed {printed}")
        canvas.drawRightString(pagesize[0] - margin, 0.7 * cm, f"Page {doc.page}")
        canvas.restoreState()

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=pagesize, leftMargin=margin, rightMargin=margin,
                            topMargin=margin, bottomMargin=1.5 * cm, title=sheet.title, author=COMPANY)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()


RENDERERS = {"csv": to_csv, "xlsx": to_xlsx, "pdf": to_pdf}


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-") or "report"


def file_response(sheet: Sheet, fmt: str, filename: str | None = None) -> HttpResponse:
    content_type, ext = FORMATS[fmt]
    body = RENDERERS[fmt](sheet)
    name = f"{slug(filename or sheet.title)}.{ext}"
    response = HttpResponse(body, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{name}"'
    response["Access-Control-Expose-Headers"] = "Content-Disposition"
    return response


def export_format(request) -> str | None:
    """``?export=csv|xlsx|pdf`` — ``format`` is DRF's own parameter."""
    from rest_framework.exceptions import ValidationError

    fmt = (request.query_params.get("export") or "").lower() or None
    if fmt and fmt not in FORMATS:
        raise ValidationError({"export": "Use csv, xlsx or pdf."})
    return fmt
