"""The Accounting page: statements from the general ledger, and their downloads.

Each test pins an identity an accountant would check by hand — the balance
sheet balances (VAT included), the trial balance agrees, a ledger account's
opening plus its movements is its closing on any date range — and that every
download carries the same figures as the screen.
"""
import csv
import io
from decimal import Decimal

from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from rest_framework.test import APIClient, APITestCase

from apps.buildings.models import Building, PropertyType, Unit, UnitClassification, UnitStatus
from apps.expenses.models import Account, Expense, ExpenseCategory, ManualIncome
from apps.payments.services import process_payment
from apps.tenants.models import Tenant

User = get_user_model()
URL = "/api/reports/accounting/"


class AccountingTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(
            username="owner", email="owner@test.com", password="testpass123!", role="owner"
        )
        cls.north = Building.objects.create(name="North Block", total_floors=1)
        cls.farm = Building.objects.create(name="Soy Farm", total_floors=1, property_type=PropertyType.FARM)
        home = Unit.objects.create(building=cls.north, label="N1", monthly_rent=Decimal("10000"),
                                   status=UnitStatus.OCCUPIED_UNPAID)
        shop = Unit.objects.create(building=cls.north, label="N2", monthly_rent=Decimal("20000"),
                                   status=UnitStatus.OCCUPIED_UNPAID,
                                   classification=UnitClassification.BUSINESS)
        cls.res = Tenant.objects.create(first_name="Res", last_name="T", id_number="R1", phone="+254700000001",
                                        unit=home, monthly_rent=Decimal("10000"), move_in_date="2026-07-01")
        cls.com = Tenant.objects.create(first_name="Com", last_name="T", id_number="C1", phone="+254700000002",
                                        unit=shop, monthly_rent=Decimal("20000"), move_in_date="2026-07-01")

        # August and September: residential 10,000; commercial 23,200 (20,000 + VAT 3,200).
        for month in (8, 9):
            process_payment(tenant=cls.res, amount=Decimal("10000"), payment_date=f"2026-{month:02d}-03",
                            period_month=month, period_year=2026)
            process_payment(tenant=cls.com, amount=Decimal("23200"), payment_date=f"2026-{month:02d}-05",
                            period_month=month, period_year=2026, reference=f"COM-{month}")
        repairs = ExpenseCategory.objects.create(name="Repairs", account=Account.objects.get(code="5200"))
        Expense.objects.create(date="2026-09-10", category=repairs, amount=Decimal("3000"), description="Roof",
                               period_month=9, period_year=2026, building=cls.north)
        ManualIncome.objects.create(date="2026-09-20", building=cls.farm, account=Account.objects.get(code="4300"),
                                    amount=Decimal("5000"), description="Soy harvest",
                                    period_month=9, period_year=2026)

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def get(self, **params):
        resp = self.client.get(URL, params)
        assert resp.status_code == 200, resp.content
        return resp.json()

    @staticmethod
    def amounts(groups):
        return {a["code"]: a["amount"] for g in groups for a in g["accounts"] if a["code"]}

    # --- statements ----------------------------------------------------

    def test_balance_sheet_balances_with_vat_owed(self):
        body = self.get(tab="balance_sheet", start="2026-09-01", end="2026-09-30")
        assert body["balanced"] is True
        # 2600 used to be left off, so any commercial rent unbalanced the sheet.
        assert self.amounts(body["liabilities"])["2600"] == 6400.0
        assert self.amounts(body["assets"])["1020"] == 10000 * 2 + 23200 * 2 + 5000 - 3000
        assert body["total_assets"] == body["total_liabilities_and_equity"]

    def test_balance_sheet_balances_for_one_building(self):
        for building in (self.north.id, self.farm.id):
            assert self.get(tab="balance_sheet", end="2026-09-30", start="2026-09-01",
                            building=building)["balanced"] is True

    def test_pnl_counts_every_income_account(self):
        body = self.get(tab="pnl", start="2026-09-01", end="2026-09-30")
        income = self.amounts(body["income"])
        assert income == {"4110": 10000.0, "4120": 20000.0, "4300": 5000.0}
        assert body["total_expenses"] == 3000.0
        assert body["net_profit"] == 32000.0

    def test_pnl_over_a_date_range(self):
        body = self.get(tab="pnl", start="2026-08-01", end="2026-09-30")
        assert body["total_income"] == 10000 * 2 + 20000 * 2 + 5000

    def test_trial_balance_agrees_in_both_modes(self):
        for mode in ("balances", "movements"):
            body = self.get(tab="trial_balance", start="2026-09-01", end="2026-09-30", mode=mode)
            assert body["balanced"] is True, mode
            assert body["total_debit"] == body["total_credit"]
        balances = self.get(tab="trial_balance", start="2026-09-01", end="2026-09-30")
        rows = {r["code"]: r for r in balances["rows"]}
        assert rows["1020"]["debit"] == 68400.0  # the bank, cumulative to 30 Sept
        assert rows["2600"]["credit"] == 6400.0

    # --- general ledger ------------------------------------------------

    def test_ledger_account_opens_moves_and_closes(self):
        body = self.get(tab="ledger", start="2026-09-01", end="2026-09-30", account="1020")
        bank = body["accounts"][0]
        assert bank["code"] == "1020"
        assert bank["opening"] == 33200.0  # August's receipts
        assert round(bank["opening"] + bank["debit"] - bank["credit"], 2) == bank["closing"]
        assert bank["lines"][-1]["balance"] == bank["closing"] == 68400.0

    def test_ledger_opening_is_the_previous_period_closing(self):
        august = self.get(tab="ledger", start="2026-08-01", end="2026-08-31", account="1020,2600")
        september = self.get(tab="ledger", start="2026-09-01", end="2026-09-30", account="1020,2600")
        closing = {a["code"]: a["closing"] for a in august["accounts"]}
        opening = {a["code"]: a["opening"] for a in september["accounts"]}
        assert closing == opening

    def test_ledger_filters_keep_the_true_balances(self):
        full = self.get(tab="ledger", start="2026-09-01", end="2026-09-30", account="1020")["accounts"][0]
        only_expenses = self.get(tab="ledger", start="2026-09-01", end="2026-09-30",
                                 account="1020", source="expense")["accounts"][0]
        assert [ln["source"] for ln in only_expenses["lines"]] == ["Expense"]
        assert only_expenses["closing"] == full["closing"]
        searched = self.get(tab="ledger", start="2026-09-01", end="2026-09-30", q="COM-9")
        assert searched["line_count"] > 0
        assert all("COM-9" in ln["reference"] for a in searched["accounts"] for ln in a["lines"])

    def test_journal_entries_each_balance(self):
        body = self.get(tab="ledger", view="journal", start="2026-09-01", end="2026-09-30")
        assert len(body["entries"]) == 4  # two receipts, an expense, the harvest
        for entry in body["entries"]:
            assert sum(ln["debit"] for ln in entry["lines"]) == sum(ln["credit"] for ln in entry["lines"])

    def test_chart_of_accounts_shows_closing_not_just_movement(self):
        body = self.get(tab="coa", start="2026-09-01", end="2026-09-30")
        bank = next(a for a in body["accounts"] if a["code"] == "1020")
        assert bank["opening"] == 33200.0
        assert bank["closing"] == 68400.0

    def test_other_tabs_answer(self):
        for tab in ("petty_cash", "budgeting"):
            self.get(tab=tab, start="2026-09-01", end="2026-09-30")

    def test_month_and_year_still_work(self):
        body = self.get(tab="pnl", month=9, year=2026)
        assert body["start"] == "2026-09-01" and body["end"] == "2026-09-30"

    def test_bad_periods_are_refused(self):
        for params in ({"start": "2026-09-30", "end": "2026-09-01"}, {"start": "yesterday", "end": "2026-09-01"},
                       {"tab": "nope"}, {"export": "doc"}):
            assert self.client.get(URL, {"tab": "pnl", **params}).status_code == 400, params

    # --- downloads -----------------------------------------------------

    def _download(self, fmt, **params):
        resp = self.client.get(URL, {"export": fmt, **params})
        assert resp.status_code == 200, resp.content[:300]
        assert "attachment;" in resp["Content-Disposition"]
        return resp

    def test_ledger_downloads_in_every_format(self):
        params = {"tab": "ledger", "start": "2026-09-01", "end": "2026-09-30", "account": "1020"}

        text = self._download("csv", **params).content.decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(text)))
        assert "General Ledger" in rows[0][0]
        assert any("01 Sep 2026 to 30 Sep 2026" in r[0] for r in rows[:4] if r)
        assert any(len(r) > 6 and r[2] == "Closing balance — 1020" and r[-1] == "68400.00" for r in rows)

        wb = load_workbook(io.BytesIO(self._download("xlsx", **params).content))
        cells = [c for row in wb.active.iter_rows(values_only=True) for c in row]
        assert "Opening balance" in cells and 68400.0 in cells

        pdf = self._download("pdf", **params).content
        assert pdf.startswith(b"%PDF")

    def test_every_tab_downloads(self):
        for tab in ("balance_sheet", "pnl", "trial_balance", "ledger", "coa", "petty_cash", "budgeting"):
            for fmt in ("csv", "xlsx", "pdf"):
                self._download(fmt, tab=tab, start="2026-09-01", end="2026-09-30")
        self._download("xlsx", tab="ledger", view="journal", start="2026-08-01", end="2026-09-30")

    def test_any_report_table_downloads(self):
        payload = {
            "title": "Rent Balances — September 2026", "subtitle": ["All buildings"],
            "columns": [{"key": "tenant", "label": "Tenant"},
                        {"key": "balance", "label": "Balance", "kind": "money"}],
            "rows": [{"tenant": 'Mercy "M" Wanjiru', "balance": 1200.5}],
            "footer": {"tenant": "Total", "balance": 1200.5},
        }
        for fmt in ("csv", "xlsx", "pdf"):
            resp = self.client.post("/api/reports/export/", {**payload, "format": fmt}, format="json")
            assert resp.status_code == 200, resp.content[:300]
        text = self.client.post("/api/reports/export/", {**payload, "format": "csv"},
                                format="json").content.decode("utf-8-sig")
        assert '"Mercy ""M"" Wanjiru",1200.50' in text
        bad = self.client.post("/api/reports/export/", {**payload, "format": "doc"}, format="json")
        assert bad.status_code == 400

    def test_downloads_need_login(self):
        assert APIClient().get(URL, {"tab": "ledger", "export": "csv"}).status_code == 401
